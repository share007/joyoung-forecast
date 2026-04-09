"""
小家电产销预测后端 API
使用 FastAPI + SQLite + LightGBM
"""
from fastapi import FastAPI, UploadFile, File, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from typing import Dict, List, Optional
from datetime import datetime
import json
import ast
import os
import sys
from io import BytesIO
import math

# 添加当前目录到路径
sys.path.insert(0, os.path.dirname(__file__))
from database import (
    get_all_products, get_product_by_name, get_sales_data_by_product_name,
    bulk_add_sales_data, save_forecast_result, get_latest_forecast,
    create_import_batch, save_planning_result, get_recent_import_batches,
    save_latest_import_file, get_latest_import_file, upsert_products_by_codes,
    begin_latest_import_file, append_latest_import_rows, finalize_latest_import_file,
    list_import_files, delete_import_file, get_all_import_rows,
    save_or_replace_daily_monthly_forecast, get_monthly_forecast_runs,
    clear_monthly_forecast_runs, upsert_material_lifecycle,
    list_material_lifecycle, get_material_lifecycle_map,
    delete_material_lifecycle_by_ids,
    save_latest_actual_sales_file, get_latest_actual_sales_file,
)
from services.forecasting import forecast_with_champion_challenger
from services.importer import parse_upload_file, infer_column_mapping, normalize_raw_df, rows_to_json_ready
from services.monthly_forecasting import forecast_next_months
from services.planning import PlanningInput, build_recommendation
import pandas as pd
import numpy as np

app = FastAPI(title="产销预测API", description="小家电企业产销预测系统")

# 允许跨域
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class ForecastRequest(BaseModel):
    """预测请求"""
    product_name: str
    forecast_days: int
    historical_data: List[dict]


class ForecastResponse(BaseModel):
    """预测响应"""
    product_name: str
    forecast_days: int
    predictions: List[dict]
    metrics: dict


class PlanningRequest(BaseModel):
    """产销建议请求"""
    product_name: str
    forecast_days: int = 30
    historical_data: Optional[List[dict]] = None
    in_transit: float = 0
    lead_time_days: int = 7
    service_level: float = 0.95
    moq: float = 1
    pack_size: float = 1


class MonthlyForecastRequest(BaseModel):
    """按产品编码-月份预测请求（默认未来3个月）"""
    forecast_months: int = 3
    column_mapping: Optional[dict] = None
    start_month: Optional[str] = None


class MaterialLifecycleItem(BaseModel):
    product_code: str
    listing_date: Optional[str] = None
    delisting_date: Optional[str] = None


class MaterialDeleteRequest(BaseModel):
    ids: List[int]


def _json_safe(value):
    if isinstance(value, dict):
        return {k: _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    if isinstance(value, pd.Timestamp):
        return value.isoformat()
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if pd.isna(value):
        return None
    return value


@app.get("/")
def root():
    return {"message": "产销预测API运行中", "version": "1.0.0"}


@app.get("/products")
def get_products():
    """获取产品列表"""
    products = get_all_products()
    return products


@app.get("/products/{product_name}/sales")
def get_product_sales(product_name: str, limit: int = None):
    """获取产品销售数据"""
    product = get_product_by_name(product_name)
    if not product:
        raise HTTPException(status_code=404, detail="产品不存在")
    
    sales_data = get_sales_data_by_product_name(product_name, limit)
    
    # 转换为前端需要的格式
    data = [{'date': row['date'], 'sales': row['sales']} for row in sales_data]
    
    return {
        "product_name": product_name,
        "data": data,
        "summary": {
            "total_sales": sum(row['sales'] for row in sales_data),
            "avg_daily": round(sum(row['sales'] for row in sales_data) / len(sales_data), 2) if sales_data else 0,
            "max_daily": max(row['sales'] for row in sales_data) if sales_data else 0,
            "min_daily": min(row['sales'] for row in sales_data) if sales_data else 0,
            "data_count": len(sales_data)
        }
    }


@app.post("/forecast")
def forecast(request: ForecastRequest):
    """执行预测"""
    if not request.historical_data:
        raise HTTPException(status_code=400, detail="需要提供历史销售数据")
    
    if request.forecast_days < 1 or request.forecast_days > 365:
        raise HTTPException(status_code=400, detail="预测天数应在1-365之间")
    
    # 检查产品是否存在
    product = get_product_by_name(request.product_name)
    if not product:
        raise HTTPException(status_code=404, detail=f"产品 '{request.product_name}' 不存在")
    
    # 执行预测（Champion-Challenger）
    output = forecast_with_champion_challenger(request.historical_data, request.forecast_days)
    result = {
        'predictions': output.predictions,
        'metrics': output.metrics,
    }
    
    # 保存预测结果到数据库
    save_forecast_result(
        product['id'],
        request.forecast_days,
        result['predictions'],
        result['metrics']
    )
    
    return ForecastResponse(
        product_name=request.product_name,
        forecast_days=request.forecast_days,
        predictions=result['predictions'],
        metrics=result['metrics']
    )


@app.post("/upload")
async def upload_file(file: UploadFile = File(...)):
    """上传销售数据文件（支持任意字段），每次导入都会生成独立记录"""
    try:
        filename = file.filename or 'uploaded_file'

        # CSV: stream in chunks for large-scale import (e.g. 1M+ rows).
        if filename.lower().endswith('.csv'):
            file.file.seek(0)

            import_file_id = None
            detected_columns = None
            columns = []
            total_rows = 0
            row_index = 0
            preview_rows = []

            for chunk in pd.read_csv(file.file, chunksize=100000):
                normalized = normalize_raw_df(chunk)
                if normalized.empty:
                    continue

                if import_file_id is None:
                    columns = normalized.columns.tolist()
                    detected_columns = infer_column_mapping(normalized)
                    import_file_id = begin_latest_import_file(
                        file_name=filename,
                        columns=columns,
                        source_type='csv',
                        detected_columns=detected_columns,
                    )

                rows = rows_to_json_ready(normalized)
                append_latest_import_rows(import_file_id, rows, start_index=row_index)
                row_index += len(rows)
                total_rows += len(rows)

                if len(preview_rows) < 30:
                    need = 30 - len(preview_rows)
                    preview_rows.extend(rows[:need])

                if detected_columns and detected_columns.get('product_code'):
                    code_col = detected_columns['product_code']
                    code_values = normalized[code_col].dropna().astype(str).tolist()
                    upsert_products_by_codes(code_values)

            if import_file_id is None:
                raise HTTPException(status_code=400, detail='CSV 文件未读取到有效数据')

            finalize_latest_import_file(import_file_id, total_rows, detected_columns=detected_columns)

            return _json_safe({
                "success": True,
                "import_file_id": import_file_id,
                "detected_columns": detected_columns,
                "sheet_name": None,
                "summary": {
                    "total_records": total_rows,
                    "columns": columns,
                    "detected_columns": detected_columns,
                    "preview_count": len(preview_rows),
                },
                "preview": preview_rows,
            })

        # Excel: keep compatibility path, but return summary + preview only.
        if filename.lower().endswith('.xlsx'):
            try:
                from openpyxl import load_workbook
            except Exception as exc:
                raise HTTPException(status_code=500, detail=f'缺少 openpyxl: {exc}')

            # Some UploadFile backends expose a spooled file object without full
            # seekable() support required by zip/openpyxl. Read bytes then wrap.
            content = await file.read()
            if not content:
                raise HTTPException(status_code=400, detail='Excel 文件内容为空')

            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            if not workbook.sheetnames:
                raise HTTPException(status_code=400, detail='Excel 文件无可用工作表')

            ws = workbook[workbook.sheetnames[0]]

            header = None
            import_file_id = None
            detected_columns = None
            total_rows = 0
            row_index = 0
            preview_rows = []
            chunk: List[dict] = []
            chunk_size = 50000

            def flush_chunk(local_chunk: List[dict]):
                nonlocal row_index, total_rows, preview_rows, detected_columns, import_file_id
                if not local_chunk:
                    return

                local_df = normalize_raw_df(pd.DataFrame(local_chunk))
                if local_df.empty:
                    return

                if import_file_id is None:
                    detected_columns = infer_column_mapping(local_df)
                    import_file_id = begin_latest_import_file(
                        file_name=filename,
                        columns=local_df.columns.tolist(),
                        sheet_name=ws.title,
                        source_type='xlsx',
                        detected_columns=detected_columns,
                    )

                rows = rows_to_json_ready(local_df)
                append_latest_import_rows(import_file_id, rows, start_index=row_index)
                row_index += len(rows)
                total_rows += len(rows)

                if len(preview_rows) < 30:
                    need = 30 - len(preview_rows)
                    preview_rows.extend(rows[:need])

                if detected_columns and detected_columns.get('product_code'):
                    code_col = detected_columns['product_code']
                    if code_col in local_df.columns:
                        upsert_products_by_codes(local_df[code_col].dropna().astype(str).tolist())

            for row in ws.iter_rows(values_only=True):
                values = list(row)
                if header is None:
                    if not any(v is not None and str(v).strip() != '' for v in values):
                        continue
                    header = [str(v).strip() if v is not None else '' for v in values]
                    continue

                if not any(v is not None and str(v).strip() != '' for v in values):
                    continue

                record = {}
                for idx, col_name in enumerate(header):
                    key = col_name or f'col_{idx + 1}'
                    value = values[idx] if idx < len(values) else None
                    record[key] = value
                chunk.append(record)

                if len(chunk) >= chunk_size:
                    flush_chunk(chunk)
                    chunk = []

            flush_chunk(chunk)

            if import_file_id is None:
                raise HTTPException(status_code=400, detail='Excel 文件未读取到有效数据')

            finalize_latest_import_file(import_file_id, total_rows, detected_columns=detected_columns)
            workbook.close()

            return _json_safe({
                "success": True,
                "import_file_id": import_file_id,
                "detected_columns": detected_columns,
                "sheet_name": ws.title,
                "summary": {
                    "total_records": total_rows,
                    "columns": header,
                    "detected_columns": detected_columns,
                    "preview_count": len(preview_rows),
                },
                "preview": preview_rows,
            })

        contents = await file.read()
        parsed = parse_upload_file(filename=filename, content=contents)
        import_file_id = save_latest_import_file(
            file_name=filename,
            rows=parsed.data,
            sheet_name=parsed.sheet_name,
            source_type='excel',
            detected_columns=parsed.detected_columns,
        )

        product_col = parsed.detected_columns.get('product_code') if parsed.detected_columns else None
        if product_col:
            product_codes = [row.get(product_col) for row in parsed.data if row.get(product_col) is not None]
            upsert_products_by_codes([str(code) for code in product_codes])

        preview_rows = parsed.data[:30]
        return _json_safe({
            "success": True,
            "import_file_id": import_file_id,
            "detected_columns": parsed.detected_columns,
            "sheet_name": parsed.sheet_name,
            "summary": {
                **parsed.summary,
                "preview_count": len(preview_rows),
            },
            "preview": preview_rows,
        })
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"文件解析失败: {str(e)}")


@app.post("/products/{product_name}/data")
def import_product_data(
    product_name: str,
    data: List[dict],
    file_name: Optional[str] = Query(default=None),
    sheet_name: Optional[str] = Query(default=None),
):
    """导入产品销售数据到数据库"""
    product = get_product_by_name(product_name)
    if not product:
        raise HTTPException(status_code=404, detail=f"产品 '{product_name}' 不存在")
    
    if not data:
        raise HTTPException(status_code=400, detail="数据不能为空")
    
    # 批量添加数据
    count = bulk_add_sales_data(product['id'], data)

    batch_id = create_import_batch(
        product_id=product['id'],
        file_name=file_name or 'manual_input',
        imported_count=count,
        sheet_name=sheet_name,
        source_type='excel' if file_name else 'manual',
        details={"records": count},
    )
    
    return {
        "success": True,
        "product_name": product_name,
        "batch_id": batch_id,
        "imported_count": count,
        "message": f"成功导入 {count} 条销售数据"
    }


@app.get("/products/{product_name}/imports")
def get_product_import_batches(product_name: str, limit: int = 10):
    """查询最近导入批次"""
    product = get_product_by_name(product_name)
    if not product:
        raise HTTPException(status_code=404, detail=f"产品 '{product_name}' 不存在")

    batches = get_recent_import_batches(product['id'], limit)
    return {
        "product_name": product_name,
        "batches": batches,
    }


@app.delete("/products/{product_name}/sales")
def delete_product_sales(product_name: str):
    """删除产品销售数据"""
    product = get_product_by_name(product_name)
    if not product:
        raise HTTPException(status_code=404, detail=f"产品 '{product_name}' 不存在")
    
    from database import delete_sales_data
    delete_sales_data(product['id'])
    
    return {"success": True, "message": f"已删除产品 '{product_name}' 的所有销售数据"}


@app.get("/forecast/history/{product_name}")
def get_forecast_history(product_name: str):
    """获取产品历史预测记录"""
    product = get_product_by_name(product_name)
    if not product:
        raise HTTPException(status_code=404, detail=f"产品 '{product_name}' 不存在")
    
    from database import get_db
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        'SELECT * FROM forecast_results WHERE product_id = ? ORDER BY created_at DESC LIMIT 10',
        (product['id'],)
    )
    rows = cursor.fetchall()
    conn.close()
    
    results = []
    for row in rows:
        predictions = _safe_parse(row['predictions'])
        metrics = _safe_parse(row['metrics'])
        results.append({
            'id': row['id'],
            'forecast_days': row['forecast_days'],
            'predictions': predictions,
            'metrics': metrics,
            'created_at': row['created_at']
        })
    
    return {"product_name": product_name, "forecasts": results}


@app.post("/planning/recommendation")
def planning_recommendation(request: PlanningRequest):
    """根据预测结果输出生产和补货建议"""
    product = get_product_by_name(request.product_name)
    if not product:
        raise HTTPException(status_code=404, detail=f"产品 '{request.product_name}' 不存在")

    if request.historical_data:
        historical_data = request.historical_data
    else:
        rows = get_sales_data_by_product_name(request.product_name)
        historical_data = [{"date": r["date"], "sales": r["sales"]} for r in rows]

    if not historical_data:
        raise HTTPException(status_code=400, detail="缺少历史销量数据")

    forecast_output = forecast_with_champion_challenger(historical_data, request.forecast_days)
    forecast_id = save_forecast_result(
        product['id'],
        request.forecast_days,
        forecast_output.predictions,
        forecast_output.metrics,
    )

    planning_input = PlanningInput(
        product_name=request.product_name,
        current_stock=float(product.get('current_stock', 0)),
        in_transit=request.in_transit,
        lead_time_days=request.lead_time_days,
        service_level=request.service_level,
        moq=request.moq,
        pack_size=request.pack_size,
    )

    recommendation = build_recommendation(
        planning_input=planning_input,
        predictions=forecast_output.predictions,
        historical_sales=historical_data,
    )
    planning_id = save_planning_result(product['id'], recommendation, forecast_id=forecast_id)

    return {
        "success": True,
        "product_name": request.product_name,
        "forecast_id": forecast_id,
        "planning_id": planning_id,
        "forecast": {
            "forecast_days": request.forecast_days,
            "predictions": forecast_output.predictions,
            "metrics": forecast_output.metrics,
        },
        "recommendation": recommendation,
    }


def _safe_parse(value):
    if value is None:
        return None
    if isinstance(value, (dict, list)):
        return value
    try:
        return json.loads(value)
    except Exception:
        try:
            return ast.literal_eval(value)
        except Exception:
            return value


@app.get('/imports/latest')
def get_latest_import():
    """获取最近一次导入的原始数据与字段信息"""
    latest = get_latest_import_file()
    if not latest:
        raise HTTPException(status_code=404, detail='尚未导入数据文件')

    return {
        'success': True,
        'file_name': latest.get('file_name'),
        'sheet_name': latest.get('sheet_name'),
        'row_count': latest.get('row_count'),
        'columns': latest.get('columns', []),
        'detected_columns': latest.get('detected_columns', {}),
        'created_at': latest.get('created_at'),
        'data': latest.get('rows', []),
    }


@app.get('/imports')
def get_imports():
    """查询全部导入记录（未删除）"""
    files = list_import_files(limit=1000)
    return {
        'success': True,
        'count': len(files),
        'files': [
            {
                'id': item.get('id'),
                'file_name': item.get('file_name'),
                'sheet_name': item.get('sheet_name'),
                'source_type': item.get('source_type'),
                'row_count': item.get('row_count', 0),
                'columns': item.get('columns', []),
                'detected_columns': item.get('detected_columns', {}),
                'created_at': item.get('created_at'),
            }
            for item in files
        ],
    }


@app.delete('/imports/{import_file_id}')
def remove_import(import_file_id: int):
    """删除导入记录，并清空该记录对应的导入数据"""
    ok = delete_import_file(import_file_id)
    if not ok:
        raise HTTPException(status_code=404, detail='导入记录不存在或已删除')
    return {
        'success': True,
        'deleted_import_file_id': import_file_id,
    }


@app.post('/forecast/monthly')
def forecast_monthly(request: MonthlyForecastRequest):
    """基于所有未删除导入数据，输出未来N个月的 产品编码-月份-销量 预测。"""
    if request.forecast_months < 1 or request.forecast_months > 12:
        raise HTTPException(status_code=400, detail='forecast_months 应在 1-12 之间')

    import_files = list_import_files(limit=1000)
    if not import_files:
        raise HTTPException(status_code=404, detail='尚未导入数据文件，请先调用 /upload')

    rows = get_all_import_rows()
    if not rows:
        raise HTTPException(status_code=400, detail='当前导入文件没有可用数据')

    try:
        lifecycle_map = get_material_lifecycle_map()
        output = forecast_next_months(
            raw_rows=rows,
            forecast_months=request.forecast_months,
            column_mapping=request.column_mapping,
            lifecycle_map=lifecycle_map,
            start_month=request.start_month,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))

    algorithm = {
        'strategy': 'Champion-Challenger（全局LightGBM vs 统计基线）',
        'selection_metric': 'WMAPE（时间顺序回测）',
        'selected_method': output.metrics.get('selected_method'),
    }

    now = datetime.now()
    run_date = f'{now.year}/{now.month}/{now.day}'
    completed_at = now.strftime('%Y-%m-%d %H:%M:%S')
    run_id = save_or_replace_daily_monthly_forecast(
        run_date=run_date,
        source_file_count=len(import_files),
        source_row_count=len(rows),
        mapping=output.mapping,
        metrics=output.metrics,
        algorithm=algorithm,
        rows=output.rows,
        completed_at=completed_at,
    )

    return _json_safe({
        'success': True,
        'run_id': run_id,
        'run_date': run_date,
        'run_completed_at': completed_at,
        'source_files': [item.get('file_name') for item in import_files],
        'source_file_count': len(import_files),
        'source_row_count': len(rows),
        'forecast_months': request.forecast_months,
        'start_month': request.start_month,
        'mapping': output.mapping,
        'algorithm': algorithm,
        'metrics': output.metrics,
        'data': output.rows,
        'output_columns': ['product_code', 'month', 'sales'],
        'interval_columns': ['lower_bound', 'upper_bound'],
    })


@app.get('/forecast/monthly/runs')
def forecast_monthly_runs(limit: int = 30):
    """获取按运行日期保留的月度预测结果（每天只保留最后一次）"""
    runs = get_monthly_forecast_runs(limit=limit)
    return _json_safe({
        'success': True,
        'count': len(runs),
        'runs': [
            {
                'id': item.get('id'),
                'run_date': item.get('run_date'),
                'created_at': item.get('created_at'),
                'source_file_count': item.get('source_file_count', 0),
                'source_row_count': item.get('source_row_count', 0),
                'mapping': item.get('mapping', {}),
                'algorithm': item.get('algorithm', {}),
                'metrics': item.get('metrics', {}),
                'data': item.get('rows', []),
            }
            for item in runs
        ],
    })


@app.delete('/forecast/monthly/runs')
def clear_forecast_monthly_runs():
    deleted = clear_monthly_forecast_runs()
    return {
        'success': True,
        'deleted_runs': deleted,
    }


@app.get('/materials')
def get_materials(limit: int = 5000, product_code: Optional[str] = None):
    items = list_material_lifecycle(limit=limit, product_code_keyword=product_code)
    return {
        'success': True,
        'count': len(items),
        'items': items,
    }


@app.post('/materials')
def upsert_materials(items: List[MaterialLifecycleItem]):
    payload = [
        {
            'product_code': item.product_code,
            'listing_date': item.listing_date,
            'delisting_date': item.delisting_date,
        }
        for item in items
    ]
    count = upsert_material_lifecycle(payload)
    return {
        'success': True,
        'upserted_count': count,
    }


@app.post('/materials/upload')
async def upload_materials(file: UploadFile = File(...)):
    filename = file.filename or 'materials_file'

    try:
        if filename.lower().endswith('.csv'):
            file.file.seek(0)
            df = pd.read_csv(file.file)
        elif filename.lower().endswith(('.xlsx', '.xls')):
            content = await file.read()
            if not content:
                raise HTTPException(status_code=400, detail='文件内容为空')
            df = pd.read_excel(BytesIO(content))
        else:
            raise HTTPException(status_code=400, detail='仅支持 CSV/XLS/XLSX')
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f'读取文件失败: {exc}')

    if df is None or df.empty:
        raise HTTPException(status_code=400, detail='文件没有可导入数据')

    normalized = {str(col).strip(): col for col in df.columns}

    def find_col(candidates: List[str]) -> Optional[str]:
        lower_map = {k.lower(): v for k, v in normalized.items()}
        for cand in candidates:
            if cand.lower() in lower_map:
                return lower_map[cand.lower()]
        for key, original in normalized.items():
            lk = key.lower()
            if any(cand.lower() in lk for cand in candidates):
                return original
        return None

    code_col = find_col(['物料编码', '产品编码', 'material_code', 'product_code', 'sku'])
    listing_col = find_col(['上市时间', '上市日期', 'listing_date', 'start_date'])
    delisting_col = find_col(['下市时间', '下市日期', 'delisting_date', 'end_date'])

    if not code_col:
        raise HTTPException(status_code=400, detail='未找到物料编码列')

    rows: List[Dict] = []
    for _, row in df.iterrows():
        code = str(row.get(code_col, '')).strip()
        if not code or code.lower() == 'nan':
            continue

        listing_raw = row.get(listing_col) if listing_col else None
        delisting_raw = row.get(delisting_col) if delisting_col else None

        listing_val = None
        delisting_val = None

        if listing_raw is not None and str(listing_raw).strip() and str(listing_raw).lower() != 'nan':
            listing_val = pd.to_datetime(listing_raw, errors='coerce')
            listing_val = listing_val.strftime('%Y-%m-%d') if pd.notna(listing_val) else None
        if delisting_raw is not None and str(delisting_raw).strip() and str(delisting_raw).lower() != 'nan':
            delisting_val = pd.to_datetime(delisting_raw, errors='coerce')
            delisting_val = delisting_val.strftime('%Y-%m-%d') if pd.notna(delisting_val) else None

        rows.append(
            {
                'product_code': code,
                'listing_date': listing_val,
                'delisting_date': delisting_val,
            }
        )

    if not rows:
        raise HTTPException(status_code=400, detail='没有可导入的物料基础数据')

    count = upsert_material_lifecycle(rows)
    return {
        'success': True,
        'file_name': filename,
        'upserted_count': count,
    }


@app.post('/materials/delete')
def delete_materials(request: MaterialDeleteRequest):
    deleted = delete_material_lifecycle_by_ids(request.ids)
    return {
        'success': True,
        'deleted_count': deleted,
    }


@app.post('/actuals/upload')
async def upload_actual_sales(file: UploadFile = File(...)):
    filename = file.filename or 'actual_sales_file'
    ext = filename.lower()

    try:
        if ext.endswith('.csv'):
            file.file.seek(0)
            df = pd.read_csv(file.file)
        elif ext.endswith(('.xlsx', '.xls')):
            content = await file.read()
            if not content:
                raise HTTPException(status_code=400, detail='文件内容为空')
            df = pd.read_excel(BytesIO(content))
        else:
            raise HTTPException(status_code=400, detail='仅支持 CSV/XLS/XLSX')
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f'读取文件失败: {exc}')

    if df is None or df.empty:
        raise HTTPException(status_code=400, detail='文件没有可导入数据')

    normalized = normalize_raw_df(df)
    detected = infer_column_mapping(normalized)
    product_col = detected.get('product_code')
    date_col = detected.get('date')
    sales_col = detected.get('sales')

    if not product_col or product_col not in normalized.columns:
        raise HTTPException(status_code=400, detail='未识别到产品编码列')
    if not date_col or date_col not in normalized.columns:
        raise HTTPException(status_code=400, detail='未识别到日期/月份列')
    if not sales_col or sales_col not in normalized.columns:
        raise HTTPException(status_code=400, detail='未识别到销量列')

    work = pd.DataFrame(
        {
            'product_code': normalized[product_col].astype(str).str.strip(),
            'date': pd.to_datetime(normalized[date_col], errors='coerce'),
            'sales': pd.to_numeric(normalized[sales_col], errors='coerce'),
        }
    )
    work = work.dropna(subset=['product_code', 'date', 'sales'])
    work = work[work['product_code'] != '']

    if work.empty:
        raise HTTPException(status_code=400, detail='没有可用于偏差计算的有效记录')

    work['month'] = work['date'].dt.to_period('M').astype(str)
    monthly = (
        work.groupby(['product_code', 'month'], as_index=False)['sales']
        .sum()
        .sort_values(['product_code', 'month'])
    )

    rows = [
        {
            'product_code': str(item['product_code']),
            'month': str(item['month']),
            'sales': float(item['sales']),
        }
        for _, item in monthly.iterrows()
    ]

    actual_file_id = save_latest_actual_sales_file(
        file_name=filename,
        rows=rows,
        columns=normalized.columns.tolist(),
        detected_columns=detected,
    )

    months = sorted({row['month'] for row in rows})
    return _json_safe(
        {
            'success': True,
            'actual_file_id': actual_file_id,
            'file_name': filename,
            'row_count': len(rows),
            'months': months,
            'detected_columns': detected,
        }
    )


@app.get('/actuals/latest')
def get_latest_actuals():
    latest = get_latest_actual_sales_file()
    if not latest:
        return {
            'success': True,
            'exists': False,
            'file': None,
        }

    return _json_safe(
        {
            'success': True,
            'exists': True,
            'file': {
                'id': latest.get('id'),
                'file_name': latest.get('file_name'),
                'row_count': latest.get('row_count', 0),
                'columns': latest.get('columns', []),
                'detected_columns': latest.get('detected_columns', {}),
                'created_at': latest.get('created_at'),
                'months': sorted({str(item.get('month')) for item in (latest.get('rows') or []) if item.get('month')}),
            },
        }
    )


@app.post('/forecast/monthly/deviation/calculate')
def calculate_monthly_deviation():
    runs = get_monthly_forecast_runs(limit=1)
    if not runs:
        raise HTTPException(status_code=404, detail='暂无预测结果，请先运行月度预测')

    latest_actual = get_latest_actual_sales_file()
    if not latest_actual:
        raise HTTPException(status_code=404, detail='暂无真实销量数据，请先导入真实销售文件')

    forecast_run = runs[0]
    forecast_rows = forecast_run.get('rows', []) or []
    actual_rows = latest_actual.get('rows', []) or []

    forecast_monthly: Dict[str, float] = {}
    for row in forecast_rows:
        month = str(row.get('month', ''))
        forecast_monthly[month] = forecast_monthly.get(month, 0.0) + float(row.get('sales', 0) or 0)

    actual_monthly: Dict[str, float] = {}
    for row in actual_rows:
        month = str(row.get('month', ''))
        actual_monthly[month] = actual_monthly.get(month, 0.0) + float(row.get('sales', 0) or 0)

    months = sorted(set(forecast_monthly.keys()) | set(actual_monthly.keys()))
    month_rows = []
    for month in months:
        forecast_val = float(forecast_monthly.get(month, 0.0))
        actual_val = float(actual_monthly.get(month, 0.0))
        diff = forecast_val - actual_val
        diff_rate = (diff / actual_val) if actual_val != 0 else None
        month_rows.append(
            {
                'month': month,
                'actual_sales': round(actual_val, 2),
                'forecast_sales': round(forecast_val, 2),
                'difference': round(diff, 2),
                'difference_rate': round(diff_rate, 6) if diff_rate is not None else None,
            }
        )

    # Detailed product-month deviation rows (for debugging/traceability)
    actual_map = {
        (str(item.get('product_code')), str(item.get('month'))): float(item.get('sales', 0) or 0)
        for item in actual_rows
    }
    detail_rows = []
    for row in forecast_rows:
        key = (str(row.get('product_code')), str(row.get('month')))
        f = float(row.get('sales', 0) or 0)
        a = float(actual_map.get(key, 0.0))
        d = f - a
        r = (d / a) if a != 0 else None
        detail_rows.append(
            {
                'product_code': key[0],
                'month': key[1],
                'actual_sales': round(a, 2),
                'forecast_sales': round(f, 2),
                'difference': round(d, 2),
                'difference_rate': round(r, 6) if r is not None else None,
            }
        )

    return _json_safe(
        {
            'success': True,
            'forecast_run': {
                'id': forecast_run.get('id'),
                'run_date': forecast_run.get('run_date'),
                'created_at': forecast_run.get('created_at'),
            },
            'actual_file': {
                'id': latest_actual.get('id'),
                'file_name': latest_actual.get('file_name'),
                'created_at': latest_actual.get('created_at'),
            },
            'months': month_rows,
            'details': detail_rows,
        }
    )


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)


STATIC_APP_DIR = os.path.join(os.path.dirname(__file__), "static_app")
if os.path.isdir(STATIC_APP_DIR):
    # Mounted last so explicit API routes keep higher priority.
    app.mount("/", StaticFiles(directory=STATIC_APP_DIR, html=True), name="frontend")